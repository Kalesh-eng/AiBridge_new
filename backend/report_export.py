"""
report_export.py — Export analytics query results to CSV, Excel, or PDF.
Called by the /report/export endpoint in main.py.
No external dependencies beyond standard library + openpyxl (already used
by the xlsx skill) + reportlab for PDF (optional, falls back to CSV if missing).
"""

import io
import csv
import json
from datetime import datetime


def export_to_csv(columns: list, rows: list, report_name: str = "report") -> bytes:
    """Export query results as CSV bytes."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(columns)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")  # utf-8-sig for Excel compatibility


def export_to_excel(columns: list, rows: list, report_name: str = "report",
                    sql: str = "", question: str = "") -> bytes:
    """Export query results as Excel (.xlsx) bytes."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        # Fallback to CSV if openpyxl not available
        return export_to_csv(columns, rows, report_name)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"

    # Header row
    header_fill  = PatternFill(start_color="185FA5", end_color="185FA5", fill_type="solid")
    header_font  = Font(color="FFFFFF", bold=True, size=11)
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill    = header_fill
        cell.font    = header_font
        cell.alignment = header_align

    # Data rows
    for row_idx, row in enumerate(rows, 2):
        for col_idx, val in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    # Auto-width columns
    for col in ws.columns:
        max_len = 0
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    # Metadata sheet
    meta = wb.create_sheet("Report Info")
    meta.append(["Report Name",  report_name])
    meta.append(["Generated At", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")])
    meta.append(["Question",     question or ""])
    meta.append(["Rows",         len(rows)])
    meta.append(["Columns",      len(columns)])
    if sql:
        meta.append(["SQL", ""])
        for line in sql.split("\n"):
            meta.append(["", line])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_to_pdf(columns: list, rows: list, report_name: str = "report",
                  sql: str = "", question: str = "") -> bytes:
    """Export query results as PDF bytes using reportlab."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
    except ImportError:
        # reportlab not installed — fall back to CSV
        return export_to_csv(columns, rows, report_name)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1*cm, rightMargin=1*cm,
                            topMargin=1*cm, bottomMargin=1*cm)

    styles = getSampleStyleSheet()
    elements = []

    # Title
    elements.append(Paragraph(f"<b>{report_name}</b>", styles['Title']))
    elements.append(Spacer(1, 0.3*cm))
    if question:
        elements.append(Paragraph(f"<i>Question: {question}</i>", styles['Normal']))
    elements.append(Paragraph(
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} · {len(rows)} rows",
        styles['Normal']
    ))
    elements.append(Spacer(1, 0.5*cm))

    # Table — max 100 rows in PDF to keep it manageable
    display_rows = rows[:100]
    table_data   = [columns] + [[str(v) if v is not None else '' for v in row]
                                 for row in display_rows]

    col_width = (landscape(A4)[0] - 2*cm) / max(len(columns), 1)
    tbl = Table(table_data, colWidths=[col_width] * len(columns), repeatRows=1)
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#185FA5')),
        ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
        ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE',   (0,0), (-1,-1), 8),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f9fafb')]),
        ('GRID',       (0,0), (-1,-1), 0.3, colors.HexColor('#e5e7eb')),
        ('VALIGN',     (0,0), (-1,-1), 'MIDDLE'),
        ('PADDING',    (0,0), (-1,-1), 4),
    ]))
    elements.append(tbl)

    if len(rows) > 100:
        elements.append(Spacer(1, 0.3*cm))
        elements.append(Paragraph(
            f"<i>Showing first 100 of {len(rows)} rows.</i>", styles['Normal']
        ))

    doc.build(elements)
    return buf.getvalue()
